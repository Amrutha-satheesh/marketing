import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_currency_sheet(wb, title, currency_symbol, divisor_from_inr, number_format_str, is_base=False):
    ws = wb.create_sheet(title=title)
    ws.views.sheetView[0].showGridLines = True
    
    PRIMARY_COLOR = "1E3A8A"      # Deep Navy
    HEADER_TEXT_COLOR = "FFFFFF"  # White
    ROW_ALT_BG = "F8FAFC"         # Light Slate
    TOTAL_ROW_BG = "E2E8F0"       # Slate
    PROFIT_GREEN_BG = "DCFCE7"    # Light Green
    PROFIT_GREEN_TXT = "166534"   # Dark Green
    MUTED_GRAY = "64748B"
    
    font_title = Font(name="Calibri", size=15, bold=True, color="1E3A8A")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color=MUTED_GRAY)
    font_header = Font(name="Calibri", size=11, bold=True, color=HEADER_TEXT_COLOR)
    font_data = Font(name="Calibri", size=10, color="0F172A")
    font_total = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_profit = Font(name="Calibri", size=10, bold=True, color=PROFIT_GREEN_TXT)
    
    fill_header = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    fill_alt = PatternFill(start_color=ROW_ALT_BG, end_color=ROW_ALT_BG, fill_type="solid")
    fill_total = PatternFill(start_color=TOTAL_ROW_BG, end_color=TOTAL_ROW_BG, fill_type="solid")
    fill_profit = PatternFill(start_color=PROFIT_GREEN_BG, end_color=PROFIT_GREEN_BG, fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="CBD5E1")
    double_border_side = Side(border_style="double", color="0F172A")
    thick_top_side = Side(border_style="medium", color="0F172A")
    
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    total_border = Border(left=thin_border_side, right=thin_border_side, top=thick_top_side, bottom=double_border_side)
    
    # Title Block
    ws.merge_cells("A1:K1")
    base_tag = " [BASE CURRENCY]" if is_base else " [CONVERTED CURRENCY]"
    ws["A1"] = f"MARKETING AGENCY - FINANCIAL TRACKER{base_tag}"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(vertical="center")
    
    ws.merge_cells("A2:K2")
    if is_base:
        ws["A2"] = f"Currency: {currency_symbol} (Primary Agency Accounting Base) | 4 Active Employees"
    else:
        ws["A2"] = f"Currency: {currency_symbol} (Converted from Base INR at rate ₹{divisor_from_inr:.2f} per 1 {currency_symbol})"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = Alignment(vertical="center")
    
    headers = [
        "Employee Name",
        f"Monthly Salary ({currency_symbol})",
        f"Cost (Rent/Extras) ({currency_symbol})",
        "Total Working Hours",
        f"Per Hour Value ({currency_symbol})",
        "Project Name",
        "Project Time (Hours)",
        f"Project Value ({currency_symbol})",
        f"Project Labor Cost ({currency_symbol})",
        f"Margin of Project ({currency_symbol})",
        "Margin %"
    ]
    
    start_row = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_num)
        cell.value = header
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border
        ws.row_dimensions[start_row].height = 28

    # Base INR values for the 4 employees
    initial_employees_inr = [
        {
            "name": "Alex Morgan (SEO Lead)",
            "inr_salary": 45000.0,
            "inr_cost": 6000.0,
            "hours": 160.0,
            "project_name": "Tech Corp SEO & Organic Growth",
            "project_hours": 80.0,
            "inr_project_value": 55000.0
        },
        {
            "name": "Sarah Jenkins (Paid Ads Manager)",
            "inr_salary": 40000.0,
            "inr_cost": 6000.0,
            "hours": 160.0,
            "project_name": "Q4 E-Commerce Meta & Google Ads",
            "project_hours": 90.0,
            "inr_project_value": 52000.0
        },
        {
            "name": "David Kim (Content Strategist)",
            "inr_salary": 35000.0,
            "inr_cost": 5000.0,
            "hours": 160.0,
            "project_name": "SaaS Lead Gen Content Hub",
            "project_hours": 70.0,
            "inr_project_value": 40000.0
        },
        {
            "name": "Elena Rostova (Creative & Social Lead)",
            "inr_salary": 38000.0,
            "inr_cost": 5000.0,
            "hours": 160.0,
            "project_name": "Fintech Brand Refresh & Social",
            "project_hours": 75.0,
            "inr_project_value": 45000.0
        }
    ]
    
    current_row = start_row + 1
    for i, emp in enumerate(initial_employees_inr):
        r = current_row
        salary_conv = round(emp["inr_salary"] / divisor_from_inr, 2)
        cost_conv = round(emp["inr_cost"] / divisor_from_inr, 2)
        proj_val_conv = round(emp["inr_project_value"] / divisor_from_inr, 2)
        
        ws.cell(row=r, column=1, value=emp["name"])
        ws.cell(row=r, column=2, value=salary_conv)
        ws.cell(row=r, column=3, value=cost_conv)
        ws.cell(row=r, column=4, value=emp["hours"])
        
        # Per hour value = Salary / Total working hours: =B{r}/D{r}
        ws.cell(row=r, column=5, value=f"=B{r}/D{r}")
        ws.cell(row=r, column=6, value=emp["project_name"])
        ws.cell(row=r, column=7, value=emp["project_hours"])
        ws.cell(row=r, column=8, value=proj_val_conv)
        
        # Project Labor Cost = Project Time * Per Hour Value
        ws.cell(row=r, column=9, value=f"=G{r}*E{r}")
        
        # Margin of project = Project Value - Project Labor Cost
        ws.cell(row=r, column=10, value=f"=H{r}-I{r}")
        
        # Margin % = Margin / Project Value
        ws.cell(row=r, column=11, value=f"=IF(H{r}>0, J{r}/H{r}, 0)")
        
        is_alt = (i % 2 == 1)
        for c in range(1, 12):
            cell = ws.cell(row=r, column=c)
            cell.font = font_data
            cell.border = cell_border
            if is_alt:
                cell.fill = fill_alt
                
        ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=2).number_format = number_format_str
        ws.cell(row=r, column=3).number_format = number_format_str
        ws.cell(row=r, column=4).number_format = "#,##0"
        ws.cell(row=r, column=5).number_format = number_format_str
        ws.cell(row=r, column=6).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=r, column=7).number_format = "#,##0.0"
        ws.cell(row=r, column=8).number_format = number_format_str
        ws.cell(row=r, column=9).number_format = number_format_str
        
        cell_margin = ws.cell(row=r, column=10)
        cell_margin.number_format = number_format_str
        cell_margin.font = font_profit
        cell_margin.fill = fill_profit
        
        cell_margin_pct = ws.cell(row=r, column=11)
        cell_margin_pct.number_format = "0.0%"
        cell_margin_pct.font = font_profit
        cell_margin_pct.fill = fill_profit
        
        ws.row_dimensions[r].height = 22
        current_row += 1

    # Totals Row (Immediately follows active employees)
    total_row = current_row
    first_data_row = start_row + 1
    last_data_row = current_row - 1
    
    ws.cell(row=total_row, column=1, value="TOTAL / SUMMARY")
    ws.cell(row=total_row, column=2, value=f"=SUM(B{first_data_row}:B{last_data_row})")
    ws.cell(row=total_row, column=3, value=f"=SUM(C{first_data_row}:C{last_data_row})")
    ws.cell(row=total_row, column=4, value=f"=SUM(D{first_data_row}:D{last_data_row})")
    ws.cell(row=total_row, column=5, value=f"=AVERAGE(E{first_data_row}:E{last_data_row})")
    ws.cell(row=total_row, column=6, value="-")
    ws.cell(row=total_row, column=7, value=f"=SUM(G{first_data_row}:G{last_data_row})")
    ws.cell(row=total_row, column=8, value=f"=SUM(H{first_data_row}:H{last_data_row})")
    ws.cell(row=total_row, column=9, value=f"=SUM(I{first_data_row}:I{last_data_row})")
    ws.cell(row=total_row, column=10, value=f"=SUM(J{first_data_row}:J{last_data_row})")
    ws.cell(row=total_row, column=11, value=f"=IF(H{total_row}>0, J{total_row}/H{total_row}, 0)")
    
    for c in range(1, 12):
        cell = ws.cell(row=total_row, column=c)
        cell.font = font_total
        cell.fill = fill_total
        cell.border = total_border
        
    ws.cell(row=total_row, column=2).number_format = number_format_str
    ws.cell(row=total_row, column=3).number_format = number_format_str
    ws.cell(row=total_row, column=4).number_format = "#,##0"
    ws.cell(row=total_row, column=5).number_format = number_format_str
    ws.cell(row=total_row, column=6).alignment = Alignment(horizontal="center")
    ws.cell(row=total_row, column=7).number_format = "#,##0.0"
    ws.cell(row=total_row, column=8).number_format = number_format_str
    ws.cell(row=total_row, column=9).number_format = number_format_str
    ws.cell(row=total_row, column=10).number_format = number_format_str
    ws.cell(row=total_row, column=11).number_format = "0.0%"
    ws.row_dimensions[total_row].height = 25

    # Column Widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = 18
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["F"].width = 32

def create_marketing_tracker():
    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet
    
    # 1. Base Sheet: INR (₹) - PRIMARY TAB
    build_currency_sheet(
        wb,
        title="INR (₹) - Base Tracker",
        currency_symbol="₹",
        divisor_from_inr=1.0,
        number_format_str="[$\u20B9-en-IN]#,##,##0.00",
        is_base=True
    )
    
    # 2. Converted Sheet: USD ($)
    build_currency_sheet(
        wb,
        title="USD ($) - Converted",
        currency_symbol="$",
        divisor_from_inr=94.94,
        number_format_str="$#,##0.00",
        is_base=False
    )
    
    # 3. Converted Sheet: AED (AED)
    build_currency_sheet(
        wb,
        title="AED - Converted",
        currency_symbol="AED ",
        divisor_from_inr=25.85,
        number_format_str='"AED "#,##0.00',
        is_base=False
    )
    
    # 4. Exchange Rates & Overhead Guide Sheet
    ws_ov = wb.create_sheet(title="Rates & Overheads")
    ws_ov.views.sheetView[0].showGridLines = True
    
    font_title = Font(name="Calibri", size=14, bold=True, color="1E3A8A")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10)
    font_data_bold = Font(name="Calibri", size=10, bold=True)
    
    ws_ov.merge_cells("A1:E1")
    ws_ov["A1"] = "CURRENCY EXCHANGE RATES & CHANGEABLE OVERHEAD EXPENSES (BASE: INR ₹)"
    ws_ov["A1"].font = font_title
    
    # Rates Table
    ws_ov.cell(row=3, column=1, value="Currency").font = font_header
    ws_ov.cell(row=3, column=1).fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    ws_ov.cell(row=3, column=2, value="Symbol").font = font_header
    ws_ov.cell(row=3, column=2).fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    ws_ov.cell(row=3, column=3, value="Conversion vs 1 Unit").font = font_header
    ws_ov.cell(row=3, column=3).fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    
    rates = [
        ("INR (Base)", "₹", "1.00 INR (Base Agency Currency)"),
        ("USD", "$", "₹94.94 INR per 1 USD (Divide INR by 94.94)"),
        ("AED", "AED", "₹25.85 INR per 1 AED (Divide INR by 25.85)")
    ]
    for idx, (code, sym, r) in enumerate(rates, start=4):
        ws_ov.cell(row=idx, column=1, value=code).font = font_data_bold
        ws_ov.cell(row=idx, column=2, value=sym).font = font_data
        ws_ov.cell(row=idx, column=3, value=r).font = font_data_bold
            
    # Overhead breakdown in 3 currencies
    ws_ov.cell(row=8, column=1, value="Expense Category").font = font_header
    ws_ov.cell(row=8, column=1).fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws_ov.cell(row=8, column=2, value="Base INR (₹)").font = font_header
    ws_ov.cell(row=8, column=2).fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws_ov.cell(row=8, column=3, value="Converted USD ($)").font = font_header
    ws_ov.cell(row=8, column=3).fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws_ov.cell(row=8, column=4, value="Converted AED").font = font_header
    ws_ov.cell(row=8, column=4).fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    
    items = [
        ("Office Rent & Co-working", 25000.0),
        ("Software Tools (SEMrush, Canva, HubSpot)", 12000.0),
        ("Utilities, Internet & Misc", 8000.0)
    ]
    
    for idx, (cat, val) in enumerate(items, start=9):
        ws_ov.cell(row=idx, column=1, value=cat).font = font_data
        
        # INR
        c1 = ws_ov.cell(row=idx, column=2, value=val)
        c1.font = font_data
        c1.number_format = "[$\u20B9-en-IN]#,##,##0.00"
        
        # USD = INR / 94.94
        c2 = ws_ov.cell(row=idx, column=3, value=f"=B{idx}/94.94")
        c2.font = font_data
        c2.number_format = "$#,##0.00"
        
        # AED = INR / 25.85
        c3 = ws_ov.cell(row=idx, column=4, value=f"=B{idx}/25.85")
        c3.font = font_data
        c3.number_format = '"AED "#,##0.00'
        
    tot_row = 12
    ws_ov.cell(row=tot_row, column=1, value="TOTAL MONTHLY OVERHEAD").font = font_data_bold
    
    c_tot1 = ws_ov.cell(row=tot_row, column=2, value="=SUM(B9:B11)")
    c_tot1.font = font_data_bold
    c_tot1.number_format = "[$\u20B9-en-IN]#,##,##0.00"
    
    c_tot2 = ws_ov.cell(row=tot_row, column=3, value="=SUM(C9:C11)")
    c_tot2.font = font_data_bold
    c_tot2.number_format = "$#,##0.00"
    
    c_tot3 = ws_ov.cell(row=tot_row, column=4, value="=SUM(D9:D11)")
    c_tot3.font = font_data_bold
    c_tot3.number_format = '"AED "#,##0.00'
    
    ws_ov.column_dimensions["A"].width = 38
    ws_ov.column_dimensions["B"].width = 20
    ws_ov.column_dimensions["C"].width = 20
    ws_ov.column_dimensions["D"].width = 20
    
    output_filename = "Marketing_Company_Financial_Tracker.xlsx"
    wb.save(output_filename)
    print(f"Spreadsheet successfully generated with INR Base as {output_filename}")

if __name__ == "__main__":
    create_marketing_tracker()
